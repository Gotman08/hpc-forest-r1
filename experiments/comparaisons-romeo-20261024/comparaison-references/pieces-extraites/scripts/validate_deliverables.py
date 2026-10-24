from pathlib import Path
import json, csv, math, re, sys, hashlib
import openpyxl
R=Path(__file__).resolve().parent;O=Path(sys.argv[1])
D=json.loads((O/'RESULTATS.json').read_text(encoding='utf8'))
assert D['summary']['pending']==0 and len(D['rows'])==62
book=openpyxl.load_workbook(O/'COMPARAISON.xlsx',data_only=True,read_only=True)
assert book.sheetnames==['Comparaison','Conditions','Mesures ROMEO','Diagnostics','Essais Forest','Moteur','Lecture']
errors=[]
for sheet in book:
    for row in sheet.iter_rows():
        for cell in row:
            if cell.data_type=='e':errors.append((sheet.title,cell.coordinate,cell.value))
assert not errors,errors
for i,r in enumerate(D['rows'],6):
    assert book['Comparaison'].cell(i,1).value==r['label']
    for c,k in [(2,'published_speedup'),(3,'reference_speedup'),(4,'forest_speedup'),(5,'forest_vs_reference')]:
        actual=book['Comparaison'].cell(i,c).value;expected=r.get(k)
        if expected is None:assert actual=='NR',(r['id'],c,actual)
        else:assert math.isclose(actual,expected,rel_tol=1e-10),(r['id'],c,actual,expected)
# This exporter omits the optional worksheet dimension metadata. Count actual
# records rather than relying on openpyxl read-only max_row (which is None).
def records(sheet):return [r for r in book[sheet].iter_rows(min_row=6,values_only=True) if r[0] is not None]
assert len(records('Mesures ROMEO'))==len(D['measurements'])
assert len(records('Diagnostics'))==19
assert len(records('Essais Forest'))==len(D['trials'])
for actual,expected in zip(records('Mesures ROMEO'),D['measurements']):
    assert list(actual[:5])==[expected[k] for k in ['group','case_id','variant','block','diagnostic']]
    if expected['seconds'] is None:assert actual[9] is None
    else:assert math.isclose(actual[9],expected['seconds']*1000,rel_tol=1e-10)
book.close()
formulas=openpyxl.load_workbook(O/'COMPARAISON.xlsx',data_only=False,read_only=False)
assert formulas['Comparaison'].freeze_panes=='B6'
assert formulas['Comparaison']['C6'].value=='=Conditions!F6/Conditions!G6'
formulas.close()
with (O/'COMPARAISON.csv').open(encoding='utf-8-sig',newline='') as f:rows=list(csv.DictReader(f))
assert [r['id'] for r in rows]==[r['id'] for r in D['rows']]
html=(O/'TABLEAU-DE-BORD.html').read_text(encoding='utf8')
embedded=re.search(r'<script id="data" type="application/json">(.*?)</script>',html,re.S)
assert embedded,'Missing embedded data'
payload=json.loads(embedded.group(1));assert payload['rows']==D['rows']
assert payload['summary']['pending']==0
scripts=re.findall(r'<script>(.*?)</script>',html,re.S)
assert len(scripts)==1
(R/'dashboard-final-check.js').write_text(scripts[0],encoding='utf8')
for target in re.findall(r'href="([^"#]+)"',html):
    if not target.startswith(('https://','http://','${')):assert (O/target).is_file(),target
assert len(list((O/'protocol').iterdir()))==22
files={p.name:{'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'bytes':p.stat().st_size} for p in sorted(O.iterdir()) if p.is_file() and p.name!='MANIFEST-LIVRABLES.json' and not p.name.endswith('.inspect.ndjson')}
(O/'MANIFEST-LIVRABLES.json').write_text(json.dumps(files,indent=2,ensure_ascii=False),encoding='utf8')
print(json.dumps({'rows':len(rows),'sheets':7,'cached_formula_errors':len(errors),'diagnostic_rows':19,'html_data_matches':True,'deliverables':list(files)},ensure_ascii=False))
